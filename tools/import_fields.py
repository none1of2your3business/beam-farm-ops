"""Import fields from the crop acres Excel sheet into the active crop year."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.database import SessionLocal, Base, engine  # noqa: E402
from app.models import AppSettings, CropYear, Field  # noqa: E402


def _num(value, default=0.0):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def infer_ownership(acres_total: float, acres_mine: float) -> str:
    if acres_mine <= 0 and acres_total > 0:
        return "custom_work"
    if acres_total > 0 and abs(acres_mine - acres_total) > 0.05:
        return "on_shares"
    return "operated_by_me"


def import_crop_acres(path: Path, year: int, dry_run: bool = False) -> int:
    wb = load_workbook(path, data_only=True)
    if "crop acres" not in wb.sheetnames:
        raise SystemExit(f"Sheet 'crop acres' not found in {path}. Sheets: {wb.sheetnames}")
    ws = wb["crop acres"]

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    created = 0
    updated = 0
    try:
        crop_year = db.scalar(select(CropYear).where(CropYear.year == year))
        if not crop_year:
            crop_year = CropYear(year=year, label=f"{year} Crop Year", is_active=1)
            db.add(crop_year)
            db.flush()
        settings = db.scalar(select(AppSettings).limit(1))
        if settings:
            settings.active_crop_year_id = crop_year.id

        # Header row is usually row 2
        for row in ws.iter_rows(min_row=3, values_only=True):
            name = row[0]
            if not name or str(name).strip() in ("", "Field", "Table 1"):
                continue
            name = str(name).strip()
            acres_total = _num(row[1])
            acres_mine = _num(row[2])
            crop = str(row[3] or "None").strip() or "None"
            if crop.lower() in ("bean", "beans", "soy", "soybean"):
                crop = "Soybeans"
            elif crop.lower() == "corn":
                crop = "Corn"
            rent = _num(row[8] if len(row) > 8 else 0)

            existing = db.scalar(
                select(Field).where(Field.crop_year_id == crop_year.id, Field.name == name)
            )
            ownership = infer_ownership(acres_total, acres_mine)
            if existing:
                field = existing
                updated += 1
            else:
                field = Field(crop_year_id=crop_year.id, name=name)
                db.add(field)
                created += 1

            field.acres_total = acres_total
            field.acres_mine = acres_mine
            field.crop = crop
            field.rent_per_acre = rent
            field.ownership_mode = ownership
            if ownership == "on_shares" and acres_total:
                field.my_share_pct = round(100 * acres_mine / acres_total, 1)
                field.lease_type = "cash_rent"
            elif ownership == "custom_work":
                field.lease_type = "none"
            else:
                field.lease_type = "cash_rent"

        if dry_run:
            db.rollback()
        else:
            db.commit()
    finally:
        db.close()

    action = "Would import" if dry_run else "Imported"
    print(f"{action}: {created} new, {updated} updated for crop year {year}")
    return created + updated


def main():
    parser = argparse.ArgumentParser(description="Import crop acres from Excel")
    parser.add_argument("xlsx", type=Path, help="Path to crop Excel file")
    parser.add_argument("--year", type=int, default=2026)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if not args.xlsx.exists():
        raise SystemExit(f"File not found: {args.xlsx}")
    import_crop_acres(args.xlsx, args.year, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
