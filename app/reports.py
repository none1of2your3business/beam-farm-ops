"""Build Excel and PDF reports for Beam Farm Ops modules."""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Callable

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    BalanceSheetItem,
    BinShare,
    CropTrial,
    Equipment,
    Field,
    GrainBin,
    GrainContract,
    GrainMovement,
    InputProduct,
    Invoice,
    Party,
    Settlement,
    TrialTreatment,
    TruckLoad,
)


REPORT_MODULES = [
    ("fields", "Fields"),
    ("bins", "Grain bins"),
    ("contracts", "Grain marketing / contracts"),
    ("parties", "Parties"),
    ("equipment", "Equipment"),
    ("invoices", "Invoices"),
    ("settlements", "Settlements"),
    ("trucking", "Trucking loads"),
    ("trials", "Crop trials"),
    ("purchases", "Products / inventory"),
    ("balance", "Balance sheet items"),
    ("movements", "Grain tickets / movements"),
    ("all", "Full farm workbook (Excel only)"),
]


def _headers_rows(db: Session, year, module: str) -> tuple[list[str], list[list[Any]]]:
    if module == "fields":
        headers = [
            "Name",
            "Crop",
            "Acres total",
            "Acres mine",
            "Ownership",
            "Rent $/ac",
            "Expected yield",
        ]
        rows = []
        if year:
            for f in db.scalars(select(Field).where(Field.crop_year_id == year.id).order_by(Field.name)):
                rows.append(
                    [
                        f.name,
                        f.crop,
                        f.acres_total,
                        f.acres_mine,
                        f.ownership_mode,
                        f.rent_per_acre,
                        f.expected_yield,
                    ]
                )
        return headers, rows

    if module == "bins":
        headers = ["Bin", "Crop", "Capacity", "Owner", "Bushels", "Bin total"]
        rows = []
        for b in db.scalars(select(GrainBin).order_by(GrainBin.crop, GrainBin.name)):
            shares = list(db.scalars(select(BinShare).where(BinShare.bin_id == b.id)))
            total = sum((s.bushels or 0) for s in shares)
            if not shares:
                rows.append([b.name, b.crop, b.capacity_bu, "", 0, total])
            for s in shares:
                rows.append([b.name, b.crop, b.capacity_bu, s.owner_name, s.bushels, total])
        return headers, rows

    if module == "contracts":
        headers = [
            "Crop",
            "Type",
            "Buyer",
            "Bushels",
            "Delivered",
            "Cash",
            "Futures",
            "Basis",
            "Month",
            "Delivery end",
            "Status",
        ]
        rows = []
        q = select(GrainContract).order_by(GrainContract.crop, GrainContract.id)
        if year:
            q = q.where(GrainContract.crop_year_id == year.id)
        for c in db.scalars(q):
            rows.append(
                [
                    c.crop,
                    c.contract_type,
                    c.buyer,
                    c.bushels,
                    c.delivered_bu,
                    c.cash_price,
                    c.futures_price,
                    c.basis,
                    c.futures_month,
                    str(c.delivery_end or ""),
                    c.status,
                ]
            )
        return headers, rows

    if module == "parties":
        headers = ["Name", "Type", "Phone", "Email", "Notes"]
        rows = [
            [p.name, p.party_type, p.phone, p.email, (p.notes or "")[:200]]
            for p in db.scalars(select(Party).order_by(Party.name))
        ]
        return headers, rows

    if module == "equipment":
        headers = ["Name", "Category", "Make", "Model", "Year", "Finance", "Market", "Loan bal", "Status"]
        rows = [
            [
                e.name,
                e.category,
                e.make,
                e.model,
                e.year,
                e.finance_status,
                e.market_value,
                e.loan_balance,
                e.status,
            ]
            for e in db.scalars(select(Equipment).order_by(Equipment.name))
        ]
        return headers, rows

    if module == "invoices":
        headers = ["Date", "Party id", "Status", "Total", "Notes"]
        rows = [
            [str(inv.invoice_date), inv.party_id, inv.status, inv.total, inv.notes]
            for inv in db.scalars(select(Invoice).order_by(Invoice.id))
        ]
        return headers, rows

    if module == "settlements":
        headers = ["Id", "Title", "Date", "Party id", "Crop year", "Status", "Total due", "Notes"]
        rows = [
            [
                s.id,
                s.title,
                str(s.settlement_date),
                s.party_id,
                s.crop_year_id,
                s.status,
                s.total_due,
                s.notes,
            ]
            for s in db.scalars(select(Settlement).order_by(Settlement.id.desc()))
        ]
        return headers, rows

    if module == "trucking":
        headers = ["Date", "Crop", "Destination", "Hauler", "Bushels", "Rate paid", "Ticket", "Notes"]
        rows = [
            [
                str(t.load_date),
                t.crop,
                t.destination,
                t.hauler,
                t.bushels,
                t.rate_paid,
                t.ticket_number,
                t.notes,
            ]
            for t in db.scalars(select(TruckLoad).order_by(TruckLoad.id.desc()).limit(500))
        ]
        return headers, rows

    if module == "trials":
        headers = ["Trial", "Crop", "Year id", "Status", "Treatments", "Question"]
        rows = []
        for t in db.scalars(select(CropTrial).order_by(CropTrial.id.desc())):
            n = len(list(db.scalars(select(TrialTreatment).where(TrialTreatment.trial_id == t.id))))
            rows.append([t.name, t.crop, t.crop_year_id, t.status, n, (t.question or "")[:80]])
        return headers, rows

    if module == "purchases":
        headers = ["Product", "Category", "Unit", "On hand", "Avg unit cost"]
        rows = [
            [p.name, p.category, p.unit, p.on_hand, p.avg_unit_cost]
            for p in db.scalars(select(InputProduct).order_by(InputProduct.name))
        ]
        return headers, rows

    if module == "balance":
        headers = ["As of", "Side", "Label", "Amount", "Current?"]
        rows = [
            [
                str(i.as_of_date),
                i.side,
                i.label,
                i.amount,
                "Y" if i.is_current else "N",
            ]
            for i in db.scalars(
                select(BalanceSheetItem).order_by(
                    BalanceSheetItem.as_of_date.desc(), BalanceSheetItem.side, BalanceSheetItem.id
                )
            )
        ]
        return headers, rows

    if module == "movements":
        headers = ["Date", "Type", "Bin id", "Field id", "Crop", "Net bu", "Moisture", "Ticket", "Owner", "Dest"]
        rows = [
            [
                str(m.move_date),
                m.move_type,
                m.bin_id,
                m.field_id,
                m.crop,
                m.net_bu,
                m.moisture,
                m.ticket_number,
                m.owner_name,
                m.destination,
            ]
            for m in db.scalars(select(GrainMovement).order_by(GrainMovement.id.desc()).limit(1000))
        ]
        return headers, rows

    return ["Info"], [["Unknown module"]]


def _sheet_title(title: str) -> str:
    """Excel sheet names: max 31 chars, no \\ / * ? : [ ]."""
    cleaned = "".join("_" if ch in r"\/*?:[]" else ch for ch in (title or "Sheet"))
    cleaned = cleaned.strip() or "Sheet"
    return cleaned[:31]


def build_excel(db: Session, year, module: str, farm_name: str) -> io.BytesIO:
    wb = Workbook()
    if module == "all":
        # multi-sheet backup
        first = True
        for key, title in REPORT_MODULES:
            if key == "all":
                continue
            headers, rows = _headers_rows(db, year, key)
            ws = wb.active if first else wb.create_sheet()
            first = False
            ws.title = _sheet_title(title)
            ws.append(headers)
            for row in rows:
                ws.append(row)
    else:
        headers, rows = _headers_rows(db, year, module)
        ws = wb.active
        label = next((t for k, t in REPORT_MODULES if k == module), module)
        ws.title = _sheet_title(label)
        ws.append([farm_name, label, date.today().isoformat()])
        ws.append([])
        ws.append(headers)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _pdf_text(value: Any) -> str:
    """Core Helvetica fonts are Latin-1 only — strip/replace unsupported chars."""
    if value is None:
        return ""
    text = str(value)
    replacements = {
        "\u2014": "-",
        "\u2013": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2022": "*",
        "\u00a0": " ",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def build_pdf(db: Session, year, module: str, farm_name: str) -> io.BytesIO:
    from fpdf import FPDF

    if module == "all":
        modules = [k for k, _ in REPORT_MODULES if k != "all"]
    else:
        modules = [module]

    pdf = FPDF(orientation="L", unit="mm", format="Letter")
    pdf.set_auto_page_break(auto=True, margin=12)

    for mod in modules:
        headers, rows = _headers_rows(db, year, mod)
        label = next((t for k, t in REPORT_MODULES if k == mod), mod)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 8, _pdf_text(f"{farm_name} - {label}"), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        year_txt = f"Crop year {year.year}" if year else "All years"
        pdf.cell(
            0,
            6,
            _pdf_text(f"{year_txt}  |  Printed {date.today().isoformat()}  |  {len(rows)} rows"),
            new_x="LMARGIN",
            new_y="NEXT",
        )
        pdf.ln(2)

        usable = pdf.w - pdf.l_margin - pdf.r_margin
        cols = max(1, len(headers))
        col_w = usable / cols
        row_h = 6

        def draw_header():
            pdf.set_font("Helvetica", "B", 8)
            pdf.set_fill_color(27, 94, 32)
            pdf.set_text_color(255, 255, 255)
            for h in headers:
                pdf.cell(col_w, row_h, _pdf_text(h)[:28], border=1, fill=True)
            pdf.ln(row_h)
            pdf.set_text_color(0, 0, 0)
            pdf.set_font("Helvetica", "", 7)

        draw_header()
        for i, row in enumerate(rows):
            if pdf.get_y() > pdf.h - 18:
                pdf.add_page()
                pdf.set_font("Helvetica", "B", 11)
                pdf.cell(0, 7, _pdf_text(f"{label} (cont.)"), new_x="LMARGIN", new_y="NEXT")
                pdf.ln(1)
                draw_header()
            fill = i % 2 == 0
            if fill:
                pdf.set_fill_color(240, 248, 240)
            for cell in row:
                pdf.cell(col_w, row_h, _pdf_text(cell)[:32], border=1, fill=fill)
            pdf.ln(row_h)

    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return buf
