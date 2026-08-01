"""Guided grain contract import (Cargill Excel / CG&B Schedules CSV) with column mapping."""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import CropYear, GrainContract

# App fields the user maps spreadsheet columns onto
CONTRACT_FIELDS: list[tuple[str, str, bool]] = [
    ("contract_number", "Contract number", True),
    ("crop", "Commodity / crop", True),
    ("bushels", "Bushels (to deliver)", True),
    ("contract_type", "Contract / price type", False),
    ("buyer", "Buyer / elevator / destination", False),
    ("futures_month", "Futures / hedge month", False),
    ("futures_price", "Futures price", False),
    ("basis", "Basis", False),
    ("cash_price", "Net / cash / scheduled price", False),
    ("delivery_start", "Shipment / delivery start", False),
    ("delivery_end", "Shipment / delivery end / due date", False),
    ("status", "Status", False),
    ("row_type", "Row type (Contract vs Adjustment)", False),
    ("crop_year_label", "Crop year label (info)", False),
]

FIELD_KEYS = [k for k, _, _ in CONTRACT_FIELDS]
FIELD_LABELS = {k: lab for k, lab, _ in CONTRACT_FIELDS}

# Normalized header → field guesses (Cargill + CG&B Schedules + common aliases)
HEADER_ALIASES: dict[str, list[str]] = {
    "contract_number": [
        "purchasecontract",
        "contractnumber",
        "contractno",
        "contractid",
        "contract#",
        "contract",
    ],
    "crop": ["commodity", "crop", "product"],
    "bushels": [
        "remaining",
        "quantitytobedelivered",
        "scheduled",
        "quantityoriginal",
        "quantity",
        "qty",
        "bushels",
        "netbu",
        "bu",
    ],
    "contract_type": [
        "contractpricetype",
        "pricetype",
        "contracttype",
        "saletype",
        "type",
    ],
    "buyer": [
        "destination",
        "entityname",
        "deliverypoint",
        "buyer",
        "elevator",
        "bizaccountname",
        "location",
    ],
    "futures_month": [
        "hedgemonth",
        "futuresmonth",
        "futmonth",
        "hedge",
    ],
    "futures_price": ["futuresprice", "futures", "futprice"],
    "basis": ["basisprice", "basis", "adjustment"],
    "cash_price": [
        "scheduledprice",
        "netprice",
        "cashprice",
        "price",
        "currentvalue",
    ],
    "delivery_start": [
        "deliverydate",
        "shipmentstart",
        "deliverystart",
        "startdate",
        "deliveryfrom",
    ],
    "delivery_end": [
        "duedate",
        "shipmentend",
        "deliveryend",
        "enddate",
        "deliveryto",
        "expirationdate",
    ],
    "status": ["status"],
    "row_type": ["rowtype"],
    "crop_year_label": ["cropyear", "year"],
}

# Exact-only aliases (avoid "contract" matching "Contract Type")
EXACT_ONLY_ALIASES = frozenset({"contract", "type", "price", "month", "futures", "basis", "bu"})

TYPE_MAP = {
    "pacer": "custom",
    "pacerrange": "custom",
    "dailyfloorplus": "min_price",
    "nobasisestablished": "basis",
    "hta": "hta",
    "basis": "basis",
    "cash": "cash",
    "forward": "forward",
    "dp": "dp",
    "delayedprice": "dp",
    "accumulator": "accumulator",
    "minprice": "min_price",
    "minimumprice": "min_price",
    "futuresonly": "futures_only",
    "future": "futures_only",
}

# CG&B Hedge Month → CME letter + calendar month
HEDGE_MONTH_CODES = {
    "jan": ("F", 1, "January"),
    "january": ("F", 1, "January"),
    "feb": ("G", 2, "February"),
    "february": ("G", 2, "February"),
    "mar": ("H", 3, "March"),
    "march": ("H", 3, "March"),
    "apr": ("J", 4, "April"),
    "april": ("J", 4, "April"),
    "may": ("K", 5, "May"),
    "jun": ("M", 6, "June"),
    "june": ("M", 6, "June"),
    "jul": ("N", 7, "July"),
    "july": ("N", 7, "July"),
    "aug": ("Q", 8, "August"),
    "august": ("Q", 8, "August"),
    "sep": ("U", 9, "September"),
    "sept": ("U", 9, "September"),
    "september": ("U", 9, "September"),
    "oct": ("V", 10, "October"),
    "october": ("V", 10, "October"),
    "nov": ("X", 11, "November"),
    "november": ("X", 11, "November"),
    "dec": ("Z", 12, "December"),
    "december": ("Z", 12, "December"),
}

CGB_MARKERS = frozenset(
    {"purchasecontract", "hedgemonth", "scheduled", "remaining", "destination"}
)


def _norm(h: str) -> str:
    return "".join(ch for ch in (h or "").lower() if ch.isalnum())


def detect_source_kind(headers: list[str]) -> str:
    """Return 'cgb', 'cargill', or 'generic' from column headers."""
    norms = {_norm(h) for h in headers if h}
    if "purchasecontract" in norms and ("hedgemonth" in norms or "scheduled" in norms):
        return "cgb"
    if "contractnumber" in norms or "rowtype" in norms or "quantitytobedelivered" in norms:
        return "cargill"
    if len(norms & CGB_MARKERS) >= 3:
        return "cgb"
    return "generic"


def _num(val: Any, default: Optional[float] = 0.0) -> Optional[float]:
    if val is None or val == "":
        return default
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", "").replace("$", "")
    if not s:
        return default
    try:
        return float(s)
    except ValueError:
        return default


def _parse_date(val: Any) -> Optional[date]:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()
    for fmt in ("%m/%d/%y", "%m/%d/%Y", "%Y-%m-%d", "%m-%d-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _crop(raw: Any) -> str:
    s = str(raw or "").strip().lower()
    if "soy" in s or s in ("beans", "bean", "sb", "sbeans"):
        return "Soybeans"
    if "corn" in s or "yellow" in s:
        return "Corn"
    if "wheat" in s:
        return "Wheat"
    return str(raw or "Corn").strip() or "Corn"


def _status(raw: Any) -> str:
    s = str(raw or "open").strip().lower()
    if "close" in s or "fill" in s or "complete" in s:
        return "closed"
    if "cancel" in s:
        return "cancelled"
    return "open"


def _contract_type(raw: Any) -> str:
    key = _norm(str(raw or ""))
    if key in TYPE_MAP:
        return TYPE_MAP[key]
    for k, v in TYPE_MAP.items():
        if k in key:
            return v
    # Keep readable label when unknown (e.g. "Cargill Managed")
    label = str(raw or "").strip()
    return label or "custom"


def normalize_futures_month(raw: Any, crop: str | None = None) -> Optional[str]:
    """Map CG&B 'NOV 2026' into app values like \"X November '26\"."""
    if raw is None or raw == "":
        return None
    s = str(raw).strip()
    if not s:
        return None
    # Already in app form
    if re.match(r"^[A-Z]\s+\w+\s+'\d{2}$", s):
        return s

    m = re.match(r"^([A-Za-z]+)\s*['/]?\s*(\d{2}|\d{4})$", s)
    if m:
        mon_key = m.group(1).lower()
        yr_raw = m.group(2)
        year = int(yr_raw) if len(yr_raw) == 4 else 2000 + int(yr_raw)
        hit = HEDGE_MONTH_CODES.get(mon_key)
        if hit:
            letter, _month_num, full = hit
            return f"{letter} {full} '{year % 100:02d}"

    # Letter + year (e.g. Z26, X'26)
    m2 = re.match(r"^([A-Za-z])\s*['/]?\s*(\d{2}|\d{4})$", s)
    if m2:
        letter = m2.group(1).upper()
        yr_raw = m2.group(2)
        year = int(yr_raw) if len(yr_raw) == 4 else 2000 + int(yr_raw)
        for _k, (code, _n, full) in HEDGE_MONTH_CODES.items():
            if code == letter:
                return f"{letter} {full} '{year % 100:02d}"
        return f"{letter} '{year % 100:02d}"

    return s


def _import_key(contract_no: str) -> str:
    return f"import:{(contract_no or '').strip()}"


def _row_type_value(row: list[Any], headers: list[str], cmap: dict) -> str:
    rt_idx = cmap.get("row_type")
    if rt_idx is not None:
        return str(_cell(row, rt_idx) or "").strip().lower()
    for hi, h in enumerate(headers):
        if _norm(h) == "rowtype":
            return str(_cell(row, hi) or "").strip().lower()
    return ""


def _alias_matches(norm_header: str, alias: str) -> bool:
    if not alias or not norm_header:
        return False
    if norm_header == alias:
        return True
    if alias in EXACT_ONLY_ALIASES:
        return False
    # Substring only when alias is distinctive
    if len(alias) >= 6 and alias in norm_header:
        return True
    return False


def guess_column_map(headers: list[str], source_kind: str | None = None) -> dict[str, int | None]:
    norms = [_norm(h) for h in headers]
    used: set[int] = set()
    out: dict[str, int | None] = {k: None for k in FIELD_KEYS}
    kind = source_kind or detect_source_kind(headers)

    prefer_order = {
        "bushels": (
            ["remaining", "scheduled", "quantitytobedelivered", "quantity", "bushels"]
            if kind == "cgb"
            else ["quantitytobedelivered", "quantityoriginal", "quantity", "bushels"]
        ),
        "buyer": (
            ["destination", "buyer", "elevator"]
            if kind == "cgb"
            else ["entityname", "deliverypoint", "buyer", "elevator", "bizaccountname"]
        ),
        "contract_number": (
            ["purchasecontract", "contractnumber", "contractno", "contractid"]
            if kind == "cgb"
            else ["contractnumber", "contractno", "contractid", "contract#", "contract"]
        ),
        "futures_month": (
            ["hedgemonth", "futuresmonth"]
            if kind == "cgb"
            else ["futuresmonth", "futmonth", "hedgemonth"]
        ),
        "cash_price": (
            ["scheduledprice", "netprice", "cashprice"]
            if kind == "cgb"
            else ["netprice", "cashprice", "price"]
        ),
        "delivery_start": (
            ["deliverydate", "shipmentstart", "deliverystart"]
            if kind == "cgb"
            else ["shipmentstart", "deliverystart", "deliverydate"]
        ),
        "delivery_end": (
            ["duedate", "expirationdate", "shipmentend"]
            if kind == "cgb"
            else ["shipmentend", "deliveryend", "duedate"]
        ),
    }

    # Pass 1: exact matches in preferred order
    for field in FIELD_KEYS:
        aliases = prefer_order.get(field, HEADER_ALIASES.get(field, []))
        for alias in aliases:
            for i, n in enumerate(norms):
                if i in used:
                    continue
                if n == alias:
                    out[field] = i
                    used.add(i)
                    break
            if out[field] is not None:
                break

    # Pass 2: fuzzy for remaining fields
    for field, aliases in HEADER_ALIASES.items():
        if out[field] is not None:
            continue
        ordered = prefer_order.get(field, aliases)
        for alias in ordered:
            for i, n in enumerate(norms):
                if i in used:
                    continue
                if _alias_matches(n, alias):
                    out[field] = i
                    used.add(i)
                    break
            if out[field] is not None:
                break
    return out


def read_tabular(path: Path, content: bytes | None = None) -> dict[str, Any]:
    """Return {sheets:[{name, headers, rows}]} for xlsx or csv."""
    name = path.name.lower()
    raw = content if content is not None else path.read_bytes()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheets = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            matrix = [list(r) for r in ws.iter_rows(values_only=True)]
            if not matrix:
                sheets.append({"name": sheet_name, "headers": [], "rows": []})
                continue
            headers = [str(c).strip() if c is not None else "" for c in matrix[0]]
            rows = []
            for row in matrix[1:]:
                cells = list(row) + [None] * max(0, len(headers) - len(row))
                rows.append(cells[: len(headers)])
            sheets.append({"name": sheet_name, "headers": headers, "rows": rows})
        return {"source_kind": "excel", "sheets": sheets}

    text = raw.decode("utf-8-sig", errors="ignore")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    matrix = [list(r) for r in reader]
    if not matrix:
        return {"source_kind": "csv", "sheets": [{"name": "Sheet1", "headers": [], "rows": []}]}
    headers = [str(c).strip() for c in matrix[0]]
    rows = [list(r) + [None] * max(0, len(headers) - len(r)) for r in matrix[1:]]
    rows = [r[: len(headers)] for r in rows]
    return {"source_kind": "csv", "sheets": [{"name": "Sheet1", "headers": headers, "rows": rows}]}


def build_payload(
    filename: str,
    stored_path: str,
    extract: dict[str, Any],
    *,
    import_source: str | None = None,
) -> dict[str, Any]:
    sheets = extract.get("sheets") or []
    sheet_index = 0
    for i, s in enumerate(sheets):
        name_l = (s.get("name") or "").lower()
        if "contract" in name_l or "schedule" in name_l:
            sheet_index = i
            break
    sheet = sheets[sheet_index] if sheets else {"name": "", "headers": [], "rows": []}
    headers = list(sheet.get("headers") or [])
    rows = list(sheet.get("rows") or [])
    detected = detect_source_kind(headers)
    source = import_source or detected
    if source in ("cgb_contracts", "cgb"):
        source = "cgb"
    if source in ("cargill_contracts", "cargill"):
        source = "cargill"
    cmap = guess_column_map(headers, source)
    default_buyer = "CG&B" if source == "cgb" else "Cargill"
    payload = {
        "step": "map",
        "stored_path": stored_path,
        "original_filename": filename,
        "source_kind": extract.get("source_kind"),
        "import_source": source,
        "default_buyer": default_buyer,
        "sheets": [
            {
                "name": s.get("name"),
                "headers": s.get("headers") or [],
                "row_count": len(s.get("rows") or []),
            }
            for s in sheets
        ],
        "sheet_index": sheet_index,
        "active_sheet": {
            "name": sheet.get("name"),
            "headers": headers,
            "rows": rows[:800],
        },
        "column_map": cmap,
        # Cargill reports have Adjustment rows; CG&B schedules do not
        "include_only_contract_rows": source != "cgb",
        "with_party_id": None,
        "proposals": [],
        "commit_summary": None,
    }
    payload["proposals"] = build_proposals(payload)
    return payload


def _cell(row: list[Any], idx: int | None) -> Any:
    if idx is None or idx < 0 or idx >= len(row):
        return None
    return row[idx]


def build_proposals(payload: dict[str, Any]) -> list[dict[str, Any]]:
    sheet = payload.get("active_sheet") or {}
    headers = list(sheet.get("headers") or [])
    rows = list(sheet.get("rows") or [])
    cmap = payload.get("column_map") or {}
    only_contract = bool(payload.get("include_only_contract_rows", True))
    default_buyer = str(payload.get("default_buyer") or "Cargill").strip() or "Cargill"
    source = payload.get("import_source") or detect_source_kind(headers)

    proposals: list[dict[str, Any]] = []
    seen_in_file: set[str] = set()

    for i, row in enumerate(rows):
        rt = _row_type_value(row, headers, cmap)
        if only_contract and rt and rt != "contract":
            continue

        cno = str(_cell(row, cmap.get("contract_number")) or "").strip()
        crop = _crop(_cell(row, cmap.get("crop")))
        bu = _num(_cell(row, cmap.get("bushels")), 0.0) or 0.0
        if not cno and bu <= 0:
            continue
        if not cno:
            cno = f"(no number · row {i + 2})"

        key = cno.lower()
        dup_in_file = key in seen_in_file
        if not dup_in_file:
            seen_in_file.add(key)

        ctype_raw = _cell(row, cmap.get("contract_type"))
        d_start = _parse_date(_cell(row, cmap.get("delivery_start")))
        d_end = _parse_date(_cell(row, cmap.get("delivery_end")))
        fut_month_raw = _cell(row, cmap.get("futures_month"))
        buyer_raw = str(_cell(row, cmap.get("buyer")) or "").strip()
        buyer = buyer_raw or default_buyer

        cash = _num(_cell(row, cmap.get("cash_price")), None)
        if cash is not None and abs(cash) < 1e-9:
            cash = None

        basis = _num(_cell(row, cmap.get("basis")), None)
        ctype = _contract_type(ctype_raw)
        if basis is not None and abs(basis) < 1e-12:
            if ctype in ("futures_only", "hta", "forward", "min_price", "custom"):
                basis = None

        proposals.append(
            {
                "row_index": i,
                "include": (not dup_in_file) and bu > 0,
                "dup_in_file": dup_in_file,
                "contract_number": cno,
                "crop": crop,
                "bushels": bu,
                "contract_type": ctype,
                "contract_type_raw": str(ctype_raw or "").strip(),
                "buyer": buyer,
                "futures_month": normalize_futures_month(fut_month_raw, crop),
                "futures_price": _num(_cell(row, cmap.get("futures_price")), None),
                "basis": basis,
                "cash_price": cash,
                "delivery_start": d_start.isoformat() if d_start else None,
                "delivery_end": d_end.isoformat() if d_end else None,
                "status": _status(_cell(row, cmap.get("status"))),
                "with_party_id": None,
                "dup_in_db": False,
                "existing_id": None,
                "source": source,
            }
        )

    return proposals


def merge_proposal_choices(
    new_proposals: list[dict[str, Any]],
    old_proposals: list[dict[str, Any]] | None,
) -> list[dict[str, Any]]:
    """Keep per-row include / with_party choices when remapping columns."""
    if not old_proposals:
        return new_proposals
    by_key = {
        (p.get("contract_number") or "").strip().lower(): p
        for p in old_proposals
    }
    for p in new_proposals:
        key = (p.get("contract_number") or "").strip().lower()
        old = by_key.get(key)
        if not old:
            continue
        if "with_party_id" in old:
            p["with_party_id"] = old.get("with_party_id")
        if not p.get("dup_in_db") and not p.get("dup_in_file") and "include" in old:
            if not old.get("dup_in_db") and not old.get("dup_in_file"):
                p["include"] = bool(old.get("include"))
    return new_proposals


def mark_db_duplicates(db: Session, year: CropYear, proposals: list[dict[str, Any]]) -> None:
    existing = list(
        db.scalars(select(GrainContract).where(GrainContract.crop_year_id == year.id))
    )
    by_key: dict[str, GrainContract] = {}
    for c in existing:
        notes = (c.notes or "").strip()
        m = re.search(r"import:([^\n|]+)", notes)
        if m:
            by_key[m.group(1).strip().lower()] = c
        if c.contract_number:
            by_key[str(c.contract_number).strip().lower()] = c

    for p in proposals:
        key = (p.get("contract_number") or "").strip().lower()
        hit = by_key.get(key)
        if hit:
            p["dup_in_db"] = True
            p["existing_id"] = hit.id
            p["include"] = False


def dump_payload(payload: dict[str, Any]) -> str:
    return json.dumps(payload, default=str)


def load_payload(raw: str | None) -> dict[str, Any]:
    if not raw:
        return {}
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return {}


def commit_proposals(
    db: Session,
    year: CropYear,
    proposals: list[dict[str, Any]],
) -> dict[str, int]:
    created = skipped_dup = skipped_off = 0
    for p in proposals:
        if not p.get("include"):
            if p.get("dup_in_db") or p.get("dup_in_file"):
                skipped_dup += 1
            else:
                skipped_off += 1
            continue
        cno = (p.get("contract_number") or "").strip()
        key = _import_key(cno)
        exists = db.scalar(
            select(GrainContract).where(
                GrainContract.crop_year_id == year.id,
                GrainContract.notes == key,
            )
        )
        if not exists and cno:
            exists = db.scalar(
                select(GrainContract).where(
                    GrainContract.crop_year_id == year.id,
                    GrainContract.contract_number == cno,
                )
            )
        if exists:
            skipped_dup += 1
            continue

        start = None
        end = None
        if p.get("delivery_start"):
            try:
                start = date.fromisoformat(p["delivery_start"])
            except ValueError:
                start = None
        if p.get("delivery_end"):
            try:
                end = date.fromisoformat(p["delivery_end"])
            except ValueError:
                end = None

        wid = p.get("with_party_id")
        if wid is not None:
            try:
                wid = int(wid)
            except (TypeError, ValueError):
                wid = None
            if wid is not None:
                from app.models import Party

                if not db.get(Party, wid):
                    wid = None

        buyer = p.get("buyer") or "Cargill"
        from app import lookups as lu

        crop_name = p.get("crop") or "Corn"
        type_name = p.get("contract_type") or "custom"
        lu.ensure(db, lu.CROP, crop_name)
        lu.ensure(db, lu.CONTRACT_TYPE, type_name)
        if buyer:
            lu.ensure(db, lu.DESTINATION, buyer)

        db.add(
            GrainContract(
                crop_year_id=year.id,
                crop=crop_name,
                contract_type=type_name,
                buyer=buyer,
                with_party_id=wid,
                bushels=float(p.get("bushels") or 0),
                delivered_bu=0.0,
                futures_price=p.get("futures_price"),
                basis=p.get("basis"),
                cash_price=p.get("cash_price"),
                futures_month=p.get("futures_month"),
                delivery_start=start,
                delivery_end=end,
                status=p.get("status") or "open",
                contract_number=cno or None,
                notes=key,
            )
        )
        created += 1
    db.commit()
    return {
        "created": created,
        "skipped_duplicates": skipped_dup,
        "skipped_unchecked": skipped_off,
    }
